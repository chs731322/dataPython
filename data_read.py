from openpyxl import load_workbook

workbook = load_workbook("students.xlsx")
worksheet = workbook["학생 점수"]

def show_students(worksheet):
    print("학생 점수 목록:")
    for row in range(2, worksheet.max_row + 1):
        name = worksheet.cell(row=row, column=1).value
        score = worksheet.cell(row=row, column=2).value
        print(f"이름: {name}, 점수: {score}점")

def find_student(worksheet, target_name=None):
    if target_name is None:
        target_name = input("검색할 학생 이름을 입력하세요: ")
    for row in range(2, worksheet.max_row + 1):
        name = worksheet.cell(row=row, column=1).value
        if name == target_name:
            return row
    return None

def add_student(workbook, worksheet, name=None, score=None):
    if name is None:
        name = input("추가할 학생 이름을 입력하세요: ")
    if name == "":
        print("학생 이름을 입력해야 합니다.")
        return
    
    # 중복 학생 확인
    row_number = find_student(worksheet, name)
    if row_number is not None:
        print(f"{name} 학생은 이미 존재합니다.")
        return
    
    # 점수가 전달되지 않았다면 사용자에게 입력 받음
    if score is None:
        score = input("추가할 학생 점수를 입력하세요: ")
    
    try:
        score = int(score)
        # 점수가 0~100 범위 내인지 확인
        if score < 0 or score > 100:
            print("점수는 0에서 100 사이의 값이어야 합니다.")
            return
    except ValueError:
        print("점수는 숫자로 입력해야 합니다.")
        return
    
    worksheet.append([name, score])
    workbook.save("students.xlsx")
    print("학생 정보가 추가되었습니다.")


while True:
    print("\n학생 정보 관리 메뉴:")
    print("1. 학생 목록 보기")
    print("2. 학생 정보 검색")
    print("3. 학생 정보 추가")
    print("4. 종료")
    choice = input("원하는 메뉴를 선택하세요: ").strip()

    if choice == "1":
        show_students(worksheet)
    elif choice == "2":
        row_number = find_student(worksheet)
        if row_number is None:
            print('해당 학생이 존재하지 않습니다.')
        else:
            name = worksheet.cell(row=row_number, column=1).value
            score = worksheet.cell(row=row_number, column=2).value
            print(f"{name} 학생의 점수는 {score}점입니다.")
    elif choice == "3":
        add_student(workbook, worksheet)
    elif choice == "4":
        print("프로그램을 종료합니다.")
        break
    else:
        print("잘못된 선택입니다. 다시 시도하세요.")