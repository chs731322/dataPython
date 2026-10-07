from openpyxl import Workbook, load_workbook
from pathlib import Path

FILE_NAME = "students.xlsx"

if not Path(FILE_NAME).exists():
    # 새로운 Excel 파일 만들기
    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "학생 점수"

    worksheet.append(["이름", "점수"])

    worksheet.append(["홍길동", 85])
    worksheet.append(["김철수", 92])
    worksheet.append(["이영희", 78])

    # Excel 파일 저장
    workbook.save(FILE_NAME)
    print("Excel 파일이 생성되었습니다.")

workbook = load_workbook(FILE_NAME)
worksheet = workbook["학생 점수"]

print("\n현재 학생 점수:")

for row in worksheet.iter_rows(min_row=2, values_only=True):
    name = row[0]
    score = row[1]
    print(f"{name} : {score}점")

# 새로운 학생 입력
while True:
    new_name = input("추가할 학생 이름을 입력하세요: ")
    
    try:
        score = int(input("추가할 학생 점수를 입력하세요: "))
        
        worksheet.append([new_name, score])
            
    except ValueError:
        print("점수는 숫자로 입력해야 합니다.")

    end = input("계속 추가하시겠습니까? (y/n): ")
    if end.lower() != 'y':
        workbook.save("students.xlsx")
        print("학생 정보가 추가되었습니다.")
        print("학생 정보 추가를 종료합니다.")
        break 