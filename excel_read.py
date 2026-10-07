from openpyxl import load_workbook

workbook = load_workbook("students.xlsx")
worksheet = workbook["학생 점수"]

print("학생 점수 목록:")

for line in worksheet.iter_rows(min_row=2, values_only=True):
    name = line[0]
    score = line[1]
    print(f"이름: {name}, 점수: {score}")